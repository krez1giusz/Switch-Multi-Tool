from __future__ import annotations

import csv
import sqlite3
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable, Optional


DATA_DIR = Path("data")
SESSIONS_DIR = Path("sessions")
REPORTS_DIR = Path("reports")
KNOWN_DB_PATH = DATA_DIR / "known_devices.sqlite3"


@dataclass
class SessionEntry:
    num: int
    timestamp: str
    port: str
    serial: str
    model: str
    ios: str
    mac: str
    profile: str
    status: str
    reason: str
    excel_status: str
    note: str
    log_path: str
    source: str = "auto"


class KnownDevicesDB:
    def __init__(self, path: Path = KNOWN_DB_PATH):
        self.path = path
        self.path.parent.mkdir(parents=True, exist_ok=True)
        self._init_db()

    def _connect(self):
        return sqlite3.connect(self.path)

    def _init_db(self):
        db = self._connect()
        try:
            db.execute(
                """
                CREATE TABLE IF NOT EXISTS devices (
                    serial TEXT PRIMARY KEY,
                    model TEXT,
                    first_seen TEXT NOT NULL,
                    last_seen TEXT NOT NULL,
                    last_status TEXT,
                    last_log TEXT,
                    test_count INTEGER NOT NULL DEFAULT 0,
                    recovery_count INTEGER NOT NULL DEFAULT 0,
                    excel_marked INTEGER NOT NULL DEFAULT 0,
                    last_session TEXT
                )
                """
            )
            db.commit()
        finally:
            db.close()

    def get(self, serial: str) -> Optional[dict]:
        serial = (serial or "").strip().upper()
        if not serial or serial == "NOT READ":
            return None
        db = self._connect()
        try:
            db.row_factory = sqlite3.Row
            row = db.execute("SELECT * FROM devices WHERE serial = ?", (serial,)).fetchone()
            return dict(row) if row else None
        finally:
            db.close()

    def upsert(self, entry: SessionEntry, session_id: str):
        serial = (entry.serial or "").strip().upper()
        if not serial or serial == "NOT READ":
            return
        now = entry.timestamp
        recovery_inc = 1 if entry.status in {"FACTORY_RESET", "WIZARD_OK"} else 0
        excel_marked = 1 if "odhaczono" in entry.excel_status.lower() or "updated" in entry.excel_status.lower() else 0
        db = self._connect()
        try:
            existing = db.execute("SELECT serial FROM devices WHERE serial = ?", (serial,)).fetchone()
            if existing:
                db.execute(
                    """
                    UPDATE devices
                    SET model = COALESCE(NULLIF(?, ''), model),
                        last_seen = ?,
                        last_status = ?,
                        last_log = ?,
                        test_count = test_count + 1,
                        recovery_count = recovery_count + ?,
                        excel_marked = CASE WHEN ? THEN 1 ELSE excel_marked END,
                        last_session = ?
                    WHERE serial = ?
                    """,
                    (entry.model, now, entry.status, entry.log_path, recovery_inc, excel_marked, session_id, serial),
                )
            else:
                db.execute(
                    """
                    INSERT INTO devices (
                        serial, model, first_seen, last_seen, last_status, last_log,
                        test_count, recovery_count, excel_marked, last_session
                    ) VALUES (?, ?, ?, ?, ?, ?, 1, ?, ?, ?)
                    """,
                    (serial, entry.model, now, now, entry.status, entry.log_path, recovery_inc, excel_marked, session_id),
                )
            db.commit()
        finally:
            db.close()


def new_session_id() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def session_dir(session_id: str) -> Path:
    path = SESSIONS_DIR / session_id
    path.mkdir(parents=True, exist_ok=True)
    return path


def append_session_csv(session_id: str, entry: SessionEntry):
    path = session_dir(session_id) / "session_results.csv"
    exists = path.exists()
    with path.open("a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(asdict(entry).keys()))
        if not exists:
            writer.writeheader()
        writer.writerow(asdict(entry))


def write_report_xlsx(session_id: str, entries: Iterable[SessionEntry], operator: str = "") -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    entries = list(entries)
    report_dir = REPORTS_DIR / session_id
    report_dir.mkdir(parents=True, exist_ok=True)
    path = report_dir / "session_report.xlsx"

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary.append(["Session", session_id])
    summary.append(["Operator", operator or "-"])
    summary.append(["Generated", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    summary.append(["Total", len(entries)])
    summary.append([])
    summary.append(["Status", "Count"])
    counts = {}
    for entry in entries:
        counts[entry.status] = counts.get(entry.status, 0) + 1
    for status, count in sorted(counts.items()):
        summary.append([status, count])

    details = wb.create_sheet("Devices")
    headers = [
        "#", "Time", "Port", "Serial", "Model", "IOS", "MAC", "Profile",
        "Status", "Reason", "Excel", "Note", "Source", "Log",
    ]
    details.append(headers)
    for entry in entries:
        details.append([
            entry.num, entry.timestamp, entry.port, entry.serial, entry.model,
            entry.ios, entry.mac, entry.profile, entry.status, entry.reason,
            entry.excel_status, entry.note, entry.source, entry.log_path,
        ])

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for sheet in (summary, details):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for column in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[get_column_letter(column[0].column)].width = min(max(max_len + 2, 10), 60)

    details.freeze_panes = "A2"
    details.auto_filter.ref = details.dimensions
    wb.save(path)
    return path
