from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import re
import shutil
from typing import Optional


GREEN_FILL = "C6EFCE"
GREEN_FONT = "006100"


@dataclass
class OrderMarkResult:
    success: bool
    message: str
    workbook_path: str = ""
    sheet_name: str = ""
    row: Optional[int] = None
    column: Optional[int] = None
    backup_path: str = ""


def create_order_workbook_backup(workbook_path: str, session_id: str) -> OrderMarkResult:
    from .app_config import APP_CONFIG

    path = Path(workbook_path)
    if not path.exists():
        return OrderMarkResult(False, f"Nie znaleziono pliku Excel: {path}", str(path))
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return OrderMarkResult(False, "Obslugiwane sa tylko pliki .xlsx i .xlsm.", str(path))

    now = datetime.now()
    backup_dir = Path(APP_CONFIG.order_excel_backup_dir) / now.strftime("%Y-%m-%d")
    backup_dir.mkdir(parents=True, exist_ok=True)
    safe_stem = re.sub(r"[^A-Za-z0-9_.-]+", "_", path.stem).strip("._") or "workbook"
    backup_path = backup_dir / f"{safe_stem}__session_{session_id}{path.suffix}"
    counter = 2
    while backup_path.exists():
        backup_path = backup_dir / f"{safe_stem}__session_{session_id}_{counter}{path.suffix}"
        counter += 1

    try:
        shutil.copy2(path, backup_path)
    except PermissionError:
        return OrderMarkResult(
            False,
            "Nie moge utworzyc backupu. Zamknij plik w Excelu i sprobuj ponownie.",
            str(path),
        )
    except Exception as exc:
        return OrderMarkResult(False, f"Nie udalo sie utworzyc backupu: {exc}", str(path))

    return OrderMarkResult(
        True,
        f"Backup oryginalnego Excela: {backup_path}",
        str(path),
        backup_path=str(backup_path),
    )


def _normalize_serial(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().upper()
    text = re.sub(r"_X[0-9A-F]{4}_", "", text)
    return "".join(ch for ch in text if ch.isalnum())


def mark_unlocked_in_workbook(
    workbook_path: str,
    serial_number: str,
    tested_at: Optional[datetime] = None,
) -> OrderMarkResult:
    try:
        from openpyxl import load_workbook
        from openpyxl.comments import Comment
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return OrderMarkResult(
            False,
            "Brak biblioteki openpyxl. Uruchom: python -m pip install -r requirements.txt",
            workbook_path,
        )

    path = Path(workbook_path)
    if not path.exists():
        return OrderMarkResult(False, f"Nie znaleziono pliku Excel: {path}", str(path))

    serial_key = _normalize_serial(serial_number)
    if not serial_key:
        return OrderMarkResult(False, "Brak serial number z testowanego switcha.", str(path))

    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return OrderMarkResult(False, "Obslugiwane sa tylko pliki .xlsx i .xlsm.", str(path))

    tested_at = tested_at or datetime.now()
    comment_line = f"Data testu: {tested_at.strftime('%Y-%m-%d %H:%M:%S')} | Status: UNLOCKED"

    try:
        workbook = load_workbook(path, keep_vba=(path.suffix.lower() == ".xlsm"))
    except PermissionError:
        return OrderMarkResult(
            False,
            "Excel blokuje plik. Zamknij arkusz i uruchom test ponownie.",
            str(path),
        )
    except Exception as exc:
        return OrderMarkResult(False, f"Nie udalo sie otworzyc Excela: {exc}", str(path))

    match = None
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if _normalize_serial(cell.value) == serial_key:
                    match = (sheet, cell)
                    break
            if match:
                break
        if match:
            break

    if not match:
        return OrderMarkResult(
            False,
            f"Serial {serial_number} nie wystepuje w pliku zamowienia.",
            str(path),
        )

    sheet, serial_cell = match
    fill = PatternFill(fill_type="solid", fgColor=GREEN_FILL)
    font = Font(color=GREEN_FONT)
    for cell in sheet[serial_cell.row]:
        cell.fill = fill
        cell.font = font

    if serial_cell.comment and serial_cell.comment.text:
        text = serial_cell.comment.text.rstrip()
        if comment_line not in text:
            text = f"{text}\n{comment_line}"
    else:
        text = comment_line
    serial_cell.comment = Comment(text, "SwitchMultiTool")

    try:
        workbook.save(path)
    except PermissionError:
        return OrderMarkResult(
            False,
            "Nie moge zapisac zmian. Zamknij plik w Excelu i sprobuj ponownie.",
            str(path),
            sheet.title,
            serial_cell.row,
            serial_cell.column,
        )
    except Exception as exc:
        return OrderMarkResult(
            False,
            f"Nie udalo sie zapisac Excela: {exc}",
            str(path),
            sheet.title,
            serial_cell.row,
            serial_cell.column,
        )

    return OrderMarkResult(
        True,
        f"Odhaczono serial {serial_number} w arkuszu {sheet.title}, wiersz {serial_cell.row}.",
        str(path),
        sheet.title,
        serial_cell.row,
        serial_cell.column,
    )


def create_sample_order_workbook(path: str) -> OrderMarkResult:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return OrderMarkResult(
            False,
            "Brak biblioteki openpyxl. Uruchom: python -m pip install -r requirements.txt",
            path,
        )

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventory"

    headers = [
        "LP",
        "Client asset ID",
        "Device type",
        "Manufacturer",
        "Model",
        "Serial Number",
        "Location",
        "Notes",
    ]
    rows = [
        [1, "INV-2026-001", "Switch", "Cisco", "WS-C2960C-8TC-L", "FOC1234A1BC", "Rack A12", "Zamowienie klienta"],
        [2, "INV-2026-002", "Switch", "Cisco", "WS-C2960C-8TC-L", "FOC1234A1BD", "Rack A12", ""],
        [3, "INV-2026-003", "Router", "Cisco", "ISR4321/K9", "FGL2201B2CD", "Rack B02", "Nie testowac F5"],
        [4, "INV-2026-004", "Switch", "Cisco", "WS-C2960X-24TS-L", "FCW1940C3DE", "Magazyn", ""],
        [5, "INV-2026-005", "Switch", "Cisco", "WS-C2960C-8TC-L", "", "Magazyn", "Przyklad bledu: brak seriala w Excelu"],
    ]

    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)

    widths = [8, 18, 14, 16, 22, 18, 16, 42]
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    try:
        workbook.save(target)
    except Exception as exc:
        return OrderMarkResult(False, f"Nie udalo sie utworzyc sample Excela: {exc}", str(target))

    return OrderMarkResult(True, f"Utworzono przykladowy plik: {target}", str(target))
